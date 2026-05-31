"""Bulk-import books from a spreadsheet (.xlsx or .csv).

Mirrors :class:`ExportService` on the read side. The flow is deliberately two
phase so the UI can show a preview before anything is written:

    rows = service.parse_file(path)   # read + validate, NO DB writes
    result = service.import_rows(rows) # insert the valid rows

`.xlsx` parsing needs the optional ``openpyxl`` package (same dependency as
export); `.csv` works on the standard library alone. Category and language are
matched by name (case-insensitive) and **auto-created** when missing, so a
spreadsheet can introduce new ones in a single pass.
"""
from __future__ import annotations

import csv
import logging
from dataclasses import dataclass
from pathlib import Path
from typing import Sequence

from ..exceptions import LibraryError
from .author_service import AuthorService
from .book_service import BookService
from .category_service import CategoryService
from .language_service import LanguageService
from .member_service import MemberService


_log = logging.getLogger(__name__)


@dataclass(frozen=True)
class ParsedRow:
    """One spreadsheet data row after mapping + validation.

    `row_num` is the 1-based line number in the source file (header = 1), so
    error messages point the user at the right row. `error` is ``None`` when
    the row is importable.
    """

    row_num: int
    title: str
    author: str
    isbn: str | None
    year: int | None
    category: str | None
    language: str | None
    total_copies: int
    error: str | None = None

    @property
    def is_valid(self) -> bool:
        return self.error is None


@dataclass(frozen=True)
class ImportResult:
    """Outcome summary returned to the UI after an import run."""

    added: int                             # rows that created a NEW book
    merged: int                            # rows folded into an existing book
    errors: list[tuple[int, str]]          # (row_num, message)
    created_categories: list[str]
    created_languages: list[str]
    created_authors: list[str]

    @property
    def failed(self) -> int:
        return len(self.errors)


@dataclass(frozen=True)
class ParsedMemberRow:
    """One members-spreadsheet data row after mapping + validation."""

    row_num: int
    name: str
    email: str | None
    phone: str | None
    error: str | None = None

    @property
    def is_valid(self) -> bool:
        return self.error is None


@dataclass(frozen=True)
class MemberImportResult:
    """Outcome summary for a members import (no related-entity creation)."""

    added: int
    errors: list[tuple[int, str]]          # (row_num, message)

    @property
    def failed(self) -> int:
        return len(self.errors)


class ImportService:
    """Read a spreadsheet of books or members and insert them via services."""

    #: Columns written to a downloadable template (also the canonical order).
    TEMPLATE_HEADERS: tuple[str, ...] = (
        "Title", "Author", "ISBN", "Year", "Category", "Language", "Copies",
    )

    #: Example rows shown in the template so the format is self-explanatory.
    _TEMPLATE_EXAMPLE: tuple[tuple, ...] = (
        ("The Pilgrim's Progress", "John Bunyan", "9780141439297", 1678,
         "Classics", "English", 2),
        ("Mere Christianity", "C. S. Lewis", "", 1952, "Theology", "English", 1),
    )

    #: Accepted header spellings → canonical field. Matched case-insensitively.
    _HEADER_ALIASES: dict[str, set[str]] = {
        "title": {"title", "book title", "name"},
        "author": {"author", "authors", "writer", "by"},
        "isbn": {"isbn", "isbn13", "isbn-13", "isbn 13", "isbn10"},
        "year": {"year", "published", "publication year", "pub year"},
        "category": {"category", "categories", "genre", "subject"},
        "language": {"language", "languages", "lang"},
        "copies": {"copies", "total copies", "quantity", "qty", "count",
                   "no of copies", "number of copies", "stock"},
    }

    #: Members template columns + examples.
    MEMBER_TEMPLATE_HEADERS: tuple[str, ...] = ("Name", "Email", "Phone")
    _MEMBER_TEMPLATE_EXAMPLE: tuple[tuple, ...] = (
        ("Alice Johnson", "alice@example.com", "9876543210"),
        ("Bob Smith", "", "9876501234"),
    )

    #: Accepted header spellings for members → canonical field.
    _MEMBER_HEADER_ALIASES: dict[str, set[str]] = {
        "name": {"name", "member", "member name", "full name"},
        "email": {"email", "e-mail", "mail", "email address"},
        "phone": {"phone", "mobile", "phone number", "mobile number",
                  "contact", "contact number", "cell"},
    }

    def __init__(
        self,
        books: BookService,
        categories: CategoryService,
        languages: LanguageService,
        authors: AuthorService,
        members: MemberService,
    ) -> None:
        self._books = books
        self._categories = categories
        self._languages = languages
        self._authors = authors
        self._members = members

    # ------------------------------------------------------------- parse  #

    def parse_file(self, path: Path) -> list[ParsedRow]:
        """Read + validate a books `path`. Raises ValueError for unusable files."""
        raw = self._read_any(path)
        if not raw:
            raise ValueError("The file is empty.")

        col = self._map_headers(raw[0], self._HEADER_ALIASES)
        if "title" not in col:
            raise ValueError(
                "The first row must contain at least a 'Title' column header."
            )

        rows: list[ParsedRow] = []
        for line_no, raw_row in enumerate(raw[1:], start=2):
            if self._is_blank(raw_row):
                continue  # skip trailing/empty rows silently
            rows.append(self._parse_row(line_no, raw_row, col))
        if not rows:
            raise ValueError("No data rows found below the header.")
        return rows

    # ------------------------------------------------------------- import #

    def import_rows(self, rows: Sequence[ParsedRow]) -> ImportResult:
        """Insert the valid rows; auto-create missing categories/languages.

        Invalid rows (and any that fail at insert time, e.g. duplicate ISBN)
        are collected into `errors` — the import is best-effort, not
        all-or-nothing, so one bad row never blocks the rest.
        """
        cat_map = {c.name.lower(): c.id for c in self._categories.list_all()}
        lang_map = {l.name.lower(): l.id for l in self._languages.list_all()}
        author_map = {a.name.lower(): a.id for a in self._authors.list_all()}
        created_cats: list[str] = []
        created_langs: list[str] = []
        created_authors: list[str] = []
        added = 0
        merged = 0
        errors: list[tuple[int, str]] = []

        for r in rows:
            if not r.is_valid:
                errors.append((r.row_num, r.error or "Invalid row"))
                continue
            try:
                author_id = self._get_or_create(
                    r.author, author_map, created_authors, self._authors
                )
                cat_id = self._get_or_create(
                    r.category, cat_map, created_cats, self._categories
                )
                lang_id = self._get_or_create(
                    r.language, lang_map, created_langs, self._languages
                )
                # Same title + author + language as an existing book → add the
                # copies to it instead of creating a duplicate record (also
                # collapses repeated rows within the same file).
                _book_id, was_merged = self._books.add_or_merge(
                    r.title, author_id, r.isbn, r.year,
                    cat_id, lang_id, r.total_copies,
                )
                if was_merged:
                    merged += 1
                else:
                    added += 1
            except LibraryError as e:
                errors.append((r.row_num, str(e)))
            except Exception as e:  # pragma: no cover - defensive
                _log.exception("Unexpected error importing row %d", r.row_num)
                errors.append((r.row_num, str(e)))

        _log.info("Import finished: %d added, %d merged, %d failed",
                  added, merged, len(errors))
        return ImportResult(added, merged, errors, created_cats, created_langs,
                            created_authors)

    # ------------------------------------------------------------ members #

    def parse_members_file(self, path: Path) -> list[ParsedMemberRow]:
        """Read + validate a members spreadsheet (.xlsx or .csv)."""
        raw = self._read_any(path)
        if not raw:
            raise ValueError("The file is empty.")
        col = self._map_headers(raw[0], self._MEMBER_HEADER_ALIASES)
        if "name" not in col:
            raise ValueError(
                "The first row must contain at least a 'Name' column header."
            )
        rows: list[ParsedMemberRow] = []
        for line_no, raw_row in enumerate(raw[1:], start=2):
            if self._is_blank(raw_row):
                continue
            name = self._cell(raw_row, col, "name")
            rows.append(ParsedMemberRow(
                row_num=line_no,
                name=name,
                email=self._cell(raw_row, col, "email") or None,
                phone=self._cell(raw_row, col, "phone") or None,
                error=None if name else "Missing Name",
            ))
        if not rows:
            raise ValueError("No data rows found below the header.")
        return rows

    def import_members(self, rows: Sequence[ParsedMemberRow]) -> MemberImportResult:
        """Insert valid member rows; best-effort with per-row error reporting."""
        added = 0
        errors: list[tuple[int, str]] = []
        for r in rows:
            if not r.is_valid:
                errors.append((r.row_num, r.error or "Invalid row"))
                continue
            try:
                self._members.add(r.name, r.email, r.phone)
                added += 1
            except LibraryError as e:
                errors.append((r.row_num, str(e)))
            except Exception as e:  # pragma: no cover - defensive
                _log.exception("Unexpected error importing member row %d", r.row_num)
                errors.append((r.row_num, str(e)))
        _log.info("Member import finished: %d added, %d failed", added, len(errors))
        return MemberImportResult(added, errors)

    # ------------------------------------------------------------ template #

    @classmethod
    def write_template(cls, path: Path) -> Path:
        """Write a starter *books* file (headers + example rows)."""
        return cls._write_template(path, cls.TEMPLATE_HEADERS,
                                   cls._TEMPLATE_EXAMPLE, sheet="Books")

    @classmethod
    def write_members_template(cls, path: Path) -> Path:
        """Write a starter *members* file (headers + example rows)."""
        return cls._write_template(path, cls.MEMBER_TEMPLATE_HEADERS,
                                   cls._MEMBER_TEMPLATE_EXAMPLE, sheet="Members")

    @classmethod
    def _write_template(
        cls, path: Path, headers: Sequence[str],
        example_rows: Sequence[Sequence], sheet: str = "Sheet1",
    ) -> Path:
        """Write headers + example rows. Returns the path actually written
        (falls back to .csv if .xlsx is asked for without openpyxl)."""
        if path.suffix.lower() == ".xlsx":
            written = cls._write_template_xlsx(path, headers, example_rows, sheet)
            if written is not None:
                return written
            path = path.with_suffix(".csv")  # degrade gracefully
        with path.open("w", newline="", encoding="utf-8") as f:
            w = csv.writer(f)
            w.writerow(headers)
            w.writerows(example_rows)
        _log.info("Import template written (CSV): %s", path)
        return path

    @classmethod
    def _write_template_xlsx(
        cls, path: Path, headers: Sequence[str],
        example_rows: Sequence[Sequence], sheet: str,
    ) -> Path | None:
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font, PatternFill
        except ImportError:
            return None
        wb = Workbook()
        ws = wb.active
        ws.title = sheet[:31]
        ws.append(list(headers))
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="4F46E5")
            cell.alignment = Alignment(horizontal="center")
        for row in example_rows:
            ws.append(list(row))
        for col_cells in ws.columns:
            letter = col_cells[0].column_letter
            width = max((len(str(c.value)) for c in col_cells if c.value), default=10)
            ws.column_dimensions[letter].width = max(12, min(40, width + 2))
        ws.freeze_panes = "A2"
        wb.save(path)
        _log.info("Import template written (XLSX): %s", path)
        return path

    # ------------------------------------------------------------ helpers #

    def _get_or_create(
        self,
        name: str | None,
        name_map: dict[str, int],
        created: list[str],
        service,
    ) -> int | None:
        if not name:
            return None
        key = name.lower()
        if key in name_map:
            return name_map[key]
        new_id = service.add(name)
        name_map[key] = new_id
        created.append(name)
        return new_id

    @staticmethod
    def _map_headers(
        header_row: Sequence, aliases: dict[str, set[str]]
    ) -> dict[str, int]:
        index: dict[str, int] = {}
        for col, raw in enumerate(header_row):
            if raw is None:
                continue
            key = str(raw).strip().lower()
            for canonical, names in aliases.items():
                if key in names and canonical not in index:
                    index[canonical] = col
        return index

    @staticmethod
    def _cell(raw: Sequence, col: dict[str, int], field: str) -> str:
        """Read a mapped column from a row as trimmed text ("" when absent)."""
        idx = col.get(field)
        if idx is None or idx >= len(raw):
            return ""
        v = raw[idx]
        if v is None:
            return ""
        # Excel hands back whole numbers as floats (e.g. 1952.0); trim the .0.
        if isinstance(v, float) and v.is_integer():
            v = int(v)
        return str(v).strip()

    def _parse_row(
        self, row_num: int, raw: Sequence, col: dict[str, int]
    ) -> ParsedRow:
        title = self._cell(raw, col, "title")
        copies = self._parse_int(self._cell(raw, col, "copies"))
        return ParsedRow(
            row_num=row_num,
            title=title,
            author=self._cell(raw, col, "author"),
            isbn=self._cell(raw, col, "isbn") or None,
            year=self._parse_int(self._cell(raw, col, "year")),
            category=self._cell(raw, col, "category") or None,
            language=self._cell(raw, col, "language") or None,
            total_copies=copies if (copies and copies >= 1) else 1,
            error=None if title else "Missing Title",
        )

    @staticmethod
    def _parse_int(text: str) -> int | None:
        text = text.strip()
        if not text:
            return None
        try:
            return int(text)
        except ValueError:
            try:
                return int(float(text))
            except ValueError:
                return None

    @staticmethod
    def _is_blank(raw_row: Sequence) -> bool:
        return all(c is None or str(c).strip() == "" for c in raw_row)

    # ------------------------------------------------------------ readers #

    @classmethod
    def _read_any(cls, path: Path) -> list[list]:
        """Read a spreadsheet to a list of rows, dispatching on the suffix."""
        suffix = path.suffix.lower()
        if suffix == ".xlsx":
            return cls._read_xlsx(path)
        if suffix in (".csv", ".txt"):
            return cls._read_csv(path)
        raise ValueError(f"Unsupported file type '{suffix}'. Use .xlsx or .csv.")

    @staticmethod
    def _read_xlsx(path: Path) -> list[list]:
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise ValueError(
                "Reading .xlsx files needs the 'openpyxl' package.\n"
                "Install it (pip install openpyxl) or save your file as .csv."
            )
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            ws = wb.active
            return [list(r) for r in ws.iter_rows(values_only=True)]
        finally:
            wb.close()

    @staticmethod
    def _read_csv(path: Path) -> list[list]:
        # utf-8-sig strips the BOM Excel writes when saving CSV.
        with path.open("r", newline="", encoding="utf-8-sig") as f:
            return [row for row in csv.reader(f)]
